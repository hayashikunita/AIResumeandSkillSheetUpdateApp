from fastapi import FastAPI, UploadFile, File, Query, Form, HTTPException, Body
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from typing import List, Optional, Tuple
import os
import json
from datetime import datetime
import re
from dotenv import load_dotenv
from docx import Document
import openpyxl

from app.ingest import extract_text
from app.parser import parse_to_schema
from app.ai_extract import extract_with_openai

BASE_DIR = os.path.dirname(__file__)
DATA_DIR = os.path.join(BASE_DIR, "data")
TEMP_DIR = os.path.join(DATA_DIR, "temp")
# Allow template both under backend/data/temp and project-root/data/temp
TEMPLATE_RESUME = os.path.join(TEMP_DIR, "temp_職務経歴書.docx")
TEMPLATE_RESUME_ALT = os.path.abspath(os.path.join(BASE_DIR, "..", "data", "temp", "temp_職務経歴書.docx"))
os.makedirs(TEMP_DIR, exist_ok=True)

load_dotenv(os.path.join(BASE_DIR, ".env"))

app = FastAPI(title="Schema Extractor")

ALLOWED_ORIGINS = [
    "http://localhost:3000",
    "http://127.0.0.1:3000",
    "*",
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def _list_temp_dirs() -> List[Tuple[str, str]]:
    """Return list of (ts, path) sorted desc."""
    if not os.path.isdir(TEMP_DIR):
        return []
    entries = []
    for name in os.listdir(TEMP_DIR):
        full = os.path.join(TEMP_DIR, name)
        if os.path.isdir(full):
            entries.append((name, full))
    entries.sort(key=lambda x: x[0], reverse=True)
    return entries


def load_schema_from_latest() -> Tuple[dict, str]:
    for ts, path in _list_temp_dirs():
        items_path = os.path.join(path, "items.json")
        if os.path.isfile(items_path):
            with open(items_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            schema = data.get("schema")
            if schema:
                return schema, ts
    raise FileNotFoundError("No schema found in data/temp")


def ensure_out_dir() -> Tuple[str, str]:
    ts = datetime.now().strftime("%Y%m%d%H%M%S")
    out_dir = os.path.join(TEMP_DIR, ts)
    os.makedirs(out_dir, exist_ok=True)
    return ts, out_dir


def render_text(schema: dict) -> str:
    lines = [
        f"期間: {schema.get('期間', '未設定')}",
        f"担当プロジェクト概要: {schema.get('担当プロジェクト概要', '未設定')}",
        f"勤務先: {schema.get('勤務先', '未設定')}",
        f"案件名: {schema.get('案件名', '未設定')}",
        "業務内容:",
    ]
    work = schema.get("業務内容") or ""
    if isinstance(work, str):
        work_lines = [ln for ln in work.splitlines() if ln.strip()]
        if work_lines:
            lines.extend(["・" + ln.strip(" ・\t") for ln in work_lines])
        else:
            lines.append("・未設定")
    else:
        lines.append("・未設定")
    lines.extend([
        f"環境: {', '.join(schema.get('環境', []) or [])}",
        f"言語: {', '.join(schema.get('言語', []) or [])}",
        f"ツール: {', '.join(schema.get('ツール', []) or [])}",
        f"フレームワーク: {', '.join(schema.get('フレームワーク', []) or [])}",
        f"ライブラリ: {', '.join(schema.get('ライブラリ', []) or [])}",
        f"規模・人数: チーム人数={schema.get('規模・人数', {}).get('チーム人数', '未設定')}, 規模={schema.get('規模・人数', {}).get('規模', '未設定')}",
        f"役割・役職: {schema.get('役割・役職', '未設定')}",
        f"担当工程: {', '.join(schema.get('担当工程：要件定義、基本設計、詳細設計、実装、単テスト、結テスト、保守運用', []) or [])}",
    ])
    return "\n".join(lines)


def _norm_text(val, default: str = "未設定") -> str:
    """Normalize various schema fields into displayable text with fallback."""
    if val is None:
        return default
    if isinstance(val, list):
        cleaned = [str(v).strip() for v in val if str(v).strip()]
        return "\n".join(cleaned) if cleaned else default
    text = str(val).strip()
    return text if text else default


def _format_period_jp(val: str) -> str:
    if not val:
        return "未設定"
    text = str(val).strip()
    # Match patterns like 2024-04〜2024-09 or 2024/4~2025/1
    m = re.match(r"^(\d{4})[/-]?(\d{1,2})(?:\s*[〜~\-]\s*(\d{4})[/-]?(\d{1,2}))?$", text)
    if m:
        y1, m1, y2, m2 = m.group(1), m.group(2), m.group(3), m.group(4)
        start = f"{y1}年{int(m1):02d}月"
        if y2 and m2:
            end = f"{y2}年{int(m2):02d}月"
            return f"{start}～{end}"
        return start
    return text


def build_resume_text(schema: dict) -> str:
    period_raw = schema.get("期間")
    period = _format_period_jp(period_raw)
    summary = _norm_text(schema.get("担当プロジェクト概要"))
    duties_raw = schema.get("業務内容")
    if isinstance(duties_raw, str):
        duties = _norm_text([ln for ln in duties_raw.splitlines()])
    else:
        duties = _norm_text(duties_raw)
    env = _norm_text(schema.get("環境"))
    langs = _norm_text(schema.get("言語"))
    tools = _norm_text(schema.get("ツール"))
    scale = schema.get("規模・人数", {})
    role = _norm_text(schema.get("役割・役職"))
    phases = _norm_text(schema.get("担当工程：要件定義、基本設計、詳細設計、実装、単テスト、結テスト、保守運用"))

    parts = [
        "【期間】",
        period,
        "",
        "【担当プロジェクト概要】",
        summary,
        "",
        "【業務内容】",
        duties,
        "",
        "【環境】",
        env,
        "",
        "【言語】",
        langs,
        "",
        "【ツール】",
        tools,
        "",
        "【規模・人数】",
        f"チーム人数：{scale.get('チーム人数', '')}",
        f"規模：{scale.get('規模', '')}",
        "",
        "【役割・役職】",
        role,
        "",
        "【担当工程：要件定義、基本設計、詳細設計、実装、単テスト、結テスト、保守運用】",
        phases,
    ]
    return "\n".join(str(p) for p in parts if p is not None)


def build_skill_text(schema: dict) -> str:
    period_raw = schema.get("期間")
    period = _format_period_jp(period_raw)
    summary = _norm_text(schema.get("担当プロジェクト概要"))
    duties_raw = schema.get("業務内容")
    if isinstance(duties_raw, str):
        duties = _norm_text([ln for ln in duties_raw.splitlines()])
    else:
        duties = _norm_text(duties_raw)
    env = _norm_text(schema.get("環境"))
    langs = _norm_text(schema.get("言語"))
    tools = _norm_text(schema.get("ツール"))
    frameworks = _norm_text(schema.get("フレームワーク"))
    libraries = _norm_text(schema.get("ライブラリ"))
    scale = schema.get("規模・人数", {})
    role = _norm_text(schema.get("役割・役職"))
    phases = _norm_text(schema.get("担当工程：要件定義、基本設計、詳細設計、実装、単テスト、結テスト、保守運用"))

    parts = [
        "【期間】",
        period,
        "",
        "【担当プロジェクト概要】",
        summary,
        "",
        "【業務内容】",
        duties,
        "",
        "【環境】",
        env,
        "",
        "【言語】",
        langs,
        "",
        "【ツール】",
        tools,
        "",
        "【フレームワーク】",
        frameworks,
        "",
        "【ライブラリ】",
        libraries,
        "",
        "【規模・人数】",
        f"チーム人数：{scale.get('チーム人数', '')}",
        f"規模：{scale.get('規模', '')}",
        "",
        "【役割・役職】",
        role,
        "",
        "【担当工程：要件定義、基本設計、詳細設計、実装、単テスト、結テスト、保守運用】",
        phases,
    ]
    return "\n".join(str(p) for p in parts if p is not None)


def save_docx(schema: dict, out_path: str):
    # If template exists (backend/data/temp or root/data/temp), fill placeholders in-place to keep layout/styles.
    template_path = TEMPLATE_RESUME if os.path.isfile(TEMPLATE_RESUME) else None
    if not template_path and os.path.isfile(TEMPLATE_RESUME_ALT):
        template_path = TEMPLATE_RESUME_ALT

    if template_path:
        doc = Document(template_path)

        project = schema.get("案件名", "案件名未設定")
        period_raw = schema.get("期間", "未設定")
        period = _format_period_jp(period_raw)
        summary = schema.get("担当プロジェクト概要", "未設定")
        company = schema.get("勤務先", "未設定")
        role = schema.get("役割・役職", "未設定")
        scale = schema.get("規模・人数", {})
        scale_text = f"チーム人数={scale.get('チーム人数', '未設定')}, 規模={scale.get('規模', '未設定')}"
        duties = _norm_text(schema.get("業務内容"))
        env = _norm_text(schema.get("環境"))
        langs = _norm_text(schema.get("言語"))
        tools = _norm_text(schema.get("ツール"))
        frameworks = _norm_text(schema.get("フレームワーク"))
        libraries = _norm_text(schema.get("ライブラリ"))
        phases = set(schema.get("担当工程：要件定義、基本設計、詳細設計、実装、単テスト、結テスト、保守運用", []) or [])

        def mark(phase: str) -> str:
            return "〇" if phase in phases else ""

        replacements = {
            "yyyy年mm月dd日": period,
            "yyyy-mm〜yyyy-mm": period,
            "name": company,
            "<name>": company,
            "ZZZ": project,
            "<ZZZ>": project,
            "職務概要": f"職務概要: {summary}",
            "<職務概要>": summary,
            "<期間>": period,
            "<案件名>": project,
            "<担当プロジェクト概要>": summary,
            "<役割・役職>": role,
            "<役割>": role,
            "<規模・人数>": scale_text,
            "<業務内容>": duties,
            "<環境>": env,
            "<言語>": langs,
            "<ツール>": tools,
            "<フレームワーク>": frameworks,
            "<ライブラリ>": libraries,
            "<要件定義>": mark("要件定義"),
            "<基本設計>": mark("基本設計"),
            "<詳細設計>": mark("詳細設計"),
            "<実装>": mark("実装"),
            "<単テスト>": mark("単テスト"),
            "<結テスト>": mark("結テスト"),
            "<保守運用>": mark("保守運用"),
        }

        def replace_text(text: str) -> str:
            out = text
            for k, v in replacements.items():
                if k in out:
                    out = out.replace(k, str(v))
            return out

        def replace_in_paragraph(paragraph):
            if not paragraph.runs:
                return
            full = "".join(run.text for run in paragraph.runs)
            replaced = replace_text(full)
            # collapse into first run to avoid split placeholders
            paragraph.runs[0].text = replaced
            for run in paragraph.runs[1:]:
                run.text = ""

        for para in doc.paragraphs:
            replace_in_paragraph(para)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        replace_in_paragraph(para)

        doc.save(out_path)
        return

    # Fallback: build a simple doc when template is missing.
    doc = Document()
    doc.add_heading(schema.get("案件名", "プロジェクト"), level=1)
    doc.add_paragraph(f"期間: {schema.get('期間', '未設定')}")
    doc.add_paragraph(f"勤務先: {schema.get('勤務先', '未設定')}")
    doc.add_paragraph(f"担当プロジェクト概要: {schema.get('担当プロジェクト概要', '未設定')}")
    doc.add_heading("業務内容", level=2)
    work = schema.get("業務内容") or ""
    if isinstance(work, str):
        for ln in [ln for ln in work.splitlines() if ln.strip()]:
            doc.add_paragraph(ln.strip(" ・\t"), style="List Bullet")
    doc.add_heading("環境 / 言語 / ツール", level=2)
    doc.add_paragraph(f"環境: {', '.join(schema.get('環境', []) or [])}")
    doc.add_paragraph(f"言語: {', '.join(schema.get('言語', []) or [])}")
    doc.add_paragraph(f"ツール: {', '.join(schema.get('ツール', []) or [])}")
    doc.add_paragraph(f"フレームワーク: {', '.join(schema.get('フレームワーク', []) or [])}")
    doc.add_paragraph(f"ライブラリ: {', '.join(schema.get('ライブラリ', []) or [])}")
    doc.add_paragraph(f"規模・人数: チーム人数={schema.get('規模・人数', {}).get('チーム人数', '未設定')}, 規模={schema.get('規模・人数', {}).get('規模', '未設定')}")
    doc.add_paragraph(f"役割・役職: {schema.get('役割・役職', '未設定')}")
    doc.add_paragraph(f"担当工程: {', '.join(schema.get('担当工程：要件定義、基本設計、詳細設計、実装、単テスト、結テスト、保守運用', []) or [])}")
    doc.save(out_path)


def save_xlsx(schema: dict, out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SkillSheet"
    rows = [
        ("スキルシート", ""),
        ("期間", schema.get("期間", "")),
        ("勤務先", schema.get("勤務先", "")),
        ("案件名", schema.get("案件名", "")),
        ("担当プロジェクト概要", schema.get("担当プロジェクト概要", "")),
        ("業務内容", schema.get("業務内容", "")),
        ("環境", "\n".join(schema.get("環境", []) or [])),
        ("言語", "\n".join(schema.get("言語", []) or [])),
        ("ツール", "\n".join(schema.get("ツール", []) or [])),
        ("フレームワーク", "\n".join(schema.get("フレームワーク", []) or [])),
        ("ライブラリ", "\n".join(schema.get("ライブラリ", []) or [])),
        ("規模・人数_チーム人数", schema.get("規模・人数", {}).get("チーム人数", "")),
        ("規模・人数_規模", schema.get("規模・人数", {}).get("規模", "")),
        ("役割・役職", schema.get("役割・役職", "")),
        ("担当工程", "\n".join(schema.get("担当工程：要件定義、基本設計、詳細設計、実装、単テスト、結テスト、保守運用", []) or [])),
    ]
    for r in rows:
        ws.append(r)
    wb.save(out_path)


def write_temp_files(schema: dict, base_dir: str = TEMP_DIR) -> dict:
    os.makedirs(base_dir, exist_ok=True)
    resume_txt_path = os.path.join(base_dir, "temp_職務経歴書_テキスト.txt")
    skill_txt_path = os.path.join(base_dir, "temp_スキルシート_テキスト.txt")
    resume_docx_path = os.path.join(base_dir, "temp_職務経歴書.docx")
    skill_xlsx_path = os.path.join(base_dir, "temp_スキルシート.xlsx")

    with open(resume_txt_path, "w", encoding="utf-8") as f:
        f.write(build_resume_text(schema))
    with open(skill_txt_path, "w", encoding="utf-8") as f:
        f.write(build_skill_text(schema))

    save_docx(schema, resume_docx_path)
    save_xlsx(schema, skill_xlsx_path)

    return {
        "resume_txt": resume_txt_path,
        "skill_txt": skill_txt_path,
        "resume_docx": resume_docx_path,
        "skill_xlsx": skill_xlsx_path,
    }


@app.get("/health")
async def health():
    return {"status": "ok"}


@app.post("/extract")
async def extract(
    save_schema_file: bool = Query(False, description="(deprecated) schema.json is always saved"),
    model: str | None = Query(None, description="Override OpenAI model name (optional)"),
    files: List[UploadFile] | None = File(None),
    manual_text: str | None = Form(None, description="Optional free-form text input"),
):
    texts: List[str] = []
    file_records: List[dict] = []

    if files:
        for f in files:
            content = await f.read()
            text = extract_text(f.filename, content)
            texts.append(text)
            file_records.append({"filename": f.filename, "raw_text": text})

    if manual_text:
        texts.append(manual_text)
        file_records.append({"filename": "manual_input.txt", "raw_text": manual_text})

    if not texts:
        raise HTTPException(status_code=400, detail="No files or manual_text provided")

    ai_used = False
    ai_error = None
    model_used = model or os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        prompt_texts = [f"# File: {rec['filename']}\n{rec['raw_text']}" for rec in file_records]
        schema = extract_with_openai(prompt_texts, model=model_used)
        ai_used = True
    except Exception as e:
        merged = "\n\n".join(texts)
        schema = parse_to_schema(merged)
        ai_error = str(e)

    result = {
        "status": "ok",
        "schema": schema,
        "ai_used": ai_used,
        "ai_error": ai_error,
        "model_used": model_used,
        "files": file_records,
    }

    # Save artifacts (schema and raw) per request
    try:
        ts = datetime.now().strftime("%Y%m%d%H%M%S")
        out_dir = os.path.join(TEMP_DIR, ts)
        os.makedirs(out_dir, exist_ok=True)
        with open(os.path.join(out_dir, "items.json"), "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        # Always persist schema.json for downstream generation.
        with open(os.path.join(out_dir, "schema.json"), "w", encoding="utf-8") as f:
            json.dump({"schema": schema}, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

    return result


@app.get("/latest-schema")
async def latest_schema():
    try:
        schema, ts = load_schema_from_latest()
        return {"status": "ok", "schema": schema, "timestamp": ts}
    except Exception as e:
        raise HTTPException(status_code=404, detail=str(e))


def _unwrap_schema(payload: Optional[dict]) -> Optional[dict]:
    if not payload:
        return payload
    if isinstance(payload, dict) and "schema" in payload and isinstance(payload.get("schema"), dict):
        return payload.get("schema")
    return payload


def _resolve_schema(schema: Optional[dict]) -> Tuple[dict, str]:
    schema = _unwrap_schema(schema)
    if schema:
        return schema, datetime.now().strftime("%Y%m%d%H%M%S")
    return load_schema_from_latest()


@app.post("/generate/text")
async def generate_text(schema: Optional[dict] = Body(None)):
    try:
        schema_use, ts_src = _resolve_schema(schema)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))
    if not schema_use or not isinstance(schema_use, dict):
        raise HTTPException(status_code=400, detail="schema is empty or invalid")
    content = build_resume_text(schema_use)
    ts, out_dir = ensure_out_dir()
    out_path = os.path.join(out_dir, "generated_text.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(content)
    return {
        "status": "ok",
        "timestamp": ts,
        "source_timestamp": ts_src,
        "text": content,
        "download_path": f"/download/{ts}/generated_text.txt",
    }


@app.post("/generate/resume")
async def generate_resume(schema: Optional[dict] = Body(None)):
    try:
        schema_use, ts_src = _resolve_schema(schema)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))
    ts, out_dir = ensure_out_dir()
    out_path = os.path.join(out_dir, "resume.docx")
    save_docx(schema_use, out_path)
    return {
        "status": "ok",
        "timestamp": ts,
        "source_timestamp": ts_src,
        "download_path": f"/download/{ts}/resume.docx",
    }


@app.post("/generate/skill-sheet")
async def generate_skill_sheet(schema: Optional[dict] = Body(None)):
    try:
        schema_use, ts_src = _resolve_schema(schema)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))
    ts, out_dir = ensure_out_dir()
    out_path = os.path.join(out_dir, "skill_sheet.xlsx")
    save_xlsx(schema_use, out_path)
    return {
        "status": "ok",
        "timestamp": ts,
        "source_timestamp": ts_src,
        "download_path": f"/download/{ts}/skill_sheet.xlsx",
    }


@app.post("/generate/temp-files")
async def generate_temp_files(schema: Optional[dict] = Body(None)):
    try:
        schema_use, ts_src = _resolve_schema(schema)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

    # Write to fixed temp paths and also snapshot to timestamped folder for download.
    root_paths = write_temp_files(schema_use, base_dir=TEMP_DIR)
    ts, out_dir = ensure_out_dir()
    snap_paths = write_temp_files(schema_use, base_dir=out_dir)

    def rel(path: str) -> str:
        return os.path.relpath(path, TEMP_DIR)

    return {
        "status": "ok",
        "timestamp": ts,
        "source_timestamp": ts_src,
        "temp_paths": root_paths,
        "download_paths": {
            "resume_txt": f"/download/{ts}/{os.path.basename(snap_paths['resume_txt'])}",
            "skill_txt": f"/download/{ts}/{os.path.basename(snap_paths['skill_txt'])}",
            "resume_docx": f"/download/{ts}/{os.path.basename(snap_paths['resume_docx'])}",
            "skill_xlsx": f"/download/{ts}/{os.path.basename(snap_paths['skill_xlsx'])}",
        },
        "relative_paths": {k: rel(v) for k, v in root_paths.items()},
    }


@app.get("/download/{ts}/{filename}")
async def download_file(ts: str, filename: str):
    path = os.path.join(TEMP_DIR, ts, filename)
    if not os.path.isfile(path):
        raise HTTPException(status_code=404, detail="file not found")
    return FileResponse(path, filename=filename)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
