import io
import os
from typing import Optional

import chardet
import openpyxl
from docx import Document
from pdfminer.high_level import extract_text as pdf_extract_text

try:
    from PIL import Image
except ImportError:
    Image = None  # Pillow optional; OCR disabled if missing

try:
    import pytesseract
except ImportError:
    pytesseract = None  # pytesseract optional; OCR disabled if missing


def _decode_bytes(data: bytes) -> str:
    if not data:
        return ""
    # try utf-8 first
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    # detect encoding
    detected = chardet.detect(data)
    enc = detected.get("encoding") or "cp932"
    try:
        return data.decode(enc, errors="ignore")
    except Exception:
        return data.decode("utf-8", errors="ignore")


def _ocr_image(content: bytes) -> str:
    if Image is None or pytesseract is None:
        return ""
    try:
        img = Image.open(io.BytesIO(content))
        return pytesseract.image_to_string(img)
    except Exception:
        return ""


def extract_text(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        try:
            return pdf_extract_text(io.BytesIO(content)) or ""
        except Exception:
            return ""
    if ext in {".docx"}:
        try:
            doc = Document(io.BytesIO(content))
            return "\n".join([p.text for p in doc.paragraphs])
        except Exception:
            return ""
    if ext in {".xlsx"}:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        parts.append(" ".join(row_vals))
            return "\n".join(parts)
        except Exception:
            return ""
    if ext in {".txt"}:
        return _decode_bytes(content)
    if ext in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}:
        return _ocr_image(content)
    # fallback: try decode as text
    return _decode_bytes(content)
